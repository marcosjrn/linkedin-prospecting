#!/usr/bin/env python3
"""Export qualified lead JSON to Excel. Requires openpyxl; no network access."""
import argparse
from collections import Counter
import json
import math
from pathlib import Path
import re
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS = {"QUALIFICADO": "QUALIFICADO",
          "NAO_QUALIFICADO": "NÃO QUALIFICADO", "NÃO QUALIFICADO": "NÃO QUALIFICADO",
          "NAO QUALIFICADO": "NÃO QUALIFICADO"}
ORDER = {"QUALIFICADO": 0, "NÃO QUALIFICADO": 1}
COLORS = {"QUALIFICADO": "E2F0D9", "NÃO QUALIFICADO": "FCE4D6"}
RESULTS = {"ATENDE": "Atende", "NAO_ATENDE": "Não atende", "NÃO ATENDE": "Não atende",
           "NAO_CONFIRMADO": "Não confirmado", "NÃO CONFIRMADO": "Não confirmado"}
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def text(value):
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(text(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def write_row(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            value = ILLEGAL.sub("", cell.value)
            if len(value) > 32767:
                raise ValueError("Texto excede limite da célula Excel; manter detalhes no JSON e resumir evidência.")
            cell.value = value
            cell.data_type = "s"  # Never evaluate scraped text as an Excel formula.
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def canonical_url(value):
    parsed = urlsplit(value)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("URL de perfil inválida")
    host = parsed.hostname.lower()
    if host not in {"linkedin.com", "www.linkedin.com"} and not host.endswith(".linkedin.com"):
        raise ValueError("URL de perfil deve ser do LinkedIn")
    if not parsed.path.startswith("/in/") or len(parsed.path.rstrip("/")) <= 4:
        raise ValueError("URL deve identificar um perfil LinkedIn")
    return "linkedin.com" + parsed.path.rstrip("/")


def range_text(company):
    bounds = company.get("employee_range") or {}
    lower, upper = bounds.get("min"), bounds.get("max")
    if lower is None and upper is None:
        return "Não informado"
    if lower is None:
        return f"Até {upper} (mínimo não informado)"
    if upper is None:
        return f"{lower}+ (máximo não informado)"
    return f"{lower}–{upper}"


def export(data, destination):
    if not isinstance(data, dict) or not isinstance(data.get("leads"), list):
        raise ValueError("Esperado objeto com metadata e lista leads")
    meta = data.get("metadata") or {}
    leads, seen = [], set()
    for original in data["leads"]:
        lead = dict(original)
        key = canonical_url(lead["person"]["linkedin_url"])
        if key in seen:
            raise ValueError("Deduplicar e unir evidências antes de exportar: " + key)
        seen.add(key)
        status = STATUS.get(lead["qualification"]["status"])
        if status is None:
            raise ValueError("Status de qualificação inválido; usar apenas QUALIFICADO ou NAO_QUALIFICADO: " + key)
        lead["display_status"] = status
        leads.append(lead)
    leads.sort(key=lambda lead: ORDER[lead["display_status"]])
    counts = Counter(lead["display_status"] for lead in leads)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leads"
    headers = ["Nome", "LinkedIn", "Cargo atual", "Empresa", "LinkedIn da empresa",
               "Localização da pessoa", "Faixa de funcionários (LinkedIn)",
               "Contagem complementar (LinkedIn)", "Status", "Justificativa",
               "Evidência no segmento", "Evidência da preferência / B2B",
               "Informações ausentes", "Alertas"]
    write_row(sheet, headers)
    detail = workbook.create_sheet("Evidências")
    write_row(detail, ["Nome", "LinkedIn", "Tipo", "Critério", "Resultado", "Campo de origem", "Evidência"])
    for lead in leads:
        person, company, qual = lead["person"], lead.get("company") or {}, lead["qualification"]
        write_row(sheet, [person.get("name"), person["linkedin_url"], person.get("current_title"),
                         company.get("name"), company.get("linkedin_url"), person.get("location"),
                         range_text(company), company.get("linkedin_employee_count"), lead["display_status"],
                         qual.get("reason"), text(qual.get("segment_evidence")),
                         text(qual.get("preference_evidence")), text(qual.get("missing_information")),
                         text(company.get("data_warnings"))])
        sheet.cell(sheet.max_row, 9).fill = PatternFill("solid", fgColor=COLORS[lead["display_status"]])
        for index in [2, 5]:
            cell = sheet.cell(sheet.max_row, index)
            if isinstance(cell.value, str) and urlsplit(cell.value).scheme in {"http", "https"}:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
        for field, label in [("required_checks", "Obrigatório"), ("preferred_checks", "Preferência")]:
            for check in qual.get(field) or []:
                write_row(detail, [person.get("name"), person["linkedin_url"], label,
                                   check.get("criterion"), RESULTS.get(check.get("result"), check.get("result")),
                                   check.get("source_field"),
                                   text(check.get("evidence"))])
    summary = workbook.create_sheet("Resumo", 1)
    write_row(summary, ["Indicador", "Valor"])
    for label, value in [("Perfis solicitados", meta.get("requested_profiles")),
                         ("Itens encontrados (antes da deduplicação)", meta.get("found_profiles")),
                         ("Pessoas únicas entregues", len(leads)),
                         ("Duplicatas removidas", meta.get("duplicates_removed")),
                         ("Empresas resolvidas", meta.get("companies_resolved"))]:
        write_row(summary, [label, value])
    for status in ORDER:
        write_row(summary, [status, counts[status]])
    for label, key in [("Custo reportado Apify (USD)", "actual_cost_usd"),
                       ("Estimativa (USD)", "estimated_cost_usd"), ("Limite autorizado (USD)", "total_budget_usd")]:
        value = meta.get(key)
        write_row(summary, [label, "Não disponível" if value is None else value])
        if isinstance(value, (int, float)):
            summary.cell(summary.max_row, 2).number_format = '"$"0.00000'
    write_row(summary, ["Limitações", text(meta.get("limitations"))])
    write_row(summary, ["Tamanho da empresa", "Faixa da página e contagem complementar do LinkedIn; não são total auditado de funcionários."])
    audit = workbook.create_sheet("Auditoria")
    write_row(audit, ["Campo", "Dados"])
    for key in ["search_filters", "actor_inputs", "provenance", "sub_budgets_usd"]:
        if key in meta:
            write_row(audit, [key, text(meta[key])])
    for lead in leads:
        if lead.get("provenance"):
            write_row(audit, [lead["person"]["linkedin_url"], text(lead["provenance"])])
    audit.sheet_state = "hidden"
    widths = {"Leads": [25, 44, 30, 28, 42, 28, 25, 25, 23, 55, 55, 55, 45, 45],
              "Resumo": [48, 85], "Evidências": [25, 44, 18, 40, 22, 35, 80], "Auditoria": [45, 90]}
    for tab in workbook:
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        tab.row_dimensions[1].height = 44
        for cell in tab[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
        for index, width in enumerate(widths[tab.title], 1):
            tab.column_dimensions[get_column_letter(index)].width = width
        for row in range(2, tab.max_row + 1):
            lines = max(sum(max(1, math.ceil(len(part) / max(10, widths[tab.title][col - 1] - 4)))
                            for part in text(tab.cell(row, col).value).split("\n"))
                        for col in range(1, tab.max_column + 1))
            tab.row_dimensions[row].height = min(409, max(30, 16 * lines + 8))
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return {"path": str(destination.resolve()), "leads": len(leads), "statuses": dict(counts)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_json", type=Path)
    parser.add_argument("output_xlsx", type=Path)
    args = parser.parse_args()
    print(json.dumps(export(json.loads(args.input_json.read_text(encoding="utf-8")), args.output_xlsx), ensure_ascii=False))
